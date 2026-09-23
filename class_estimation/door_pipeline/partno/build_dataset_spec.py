"""현장 수집 도어 데이터 명세서 생성 — DB·품번표·검사 결과에서 수치를 읽어 DOC_dataset_spec.md + .docx 를 만든다 (재생성 가능).

  python partno/build_dataset_spec.py            # → ../DOC_dataset_spec.md, ../DOC_dataset_spec.docx
  python partno/build_dataset_spec.py --no-docx

수치 기준: 실행 시점 DB(door_pipeline.db). 옵션(레이더) 수량은 확정값(user) 과 자동(auto) 을 구분해 적는다.
"""
import argparse, collections, json, os, re, sqlite3, sys, time

HERE = os.path.dirname(os.path.abspath(__file__)); DOOR = os.path.dirname(HERE)
ROOT = os.path.normpath(os.path.join(DOOR, '..', '..'))
sys.path.insert(0, DOOR); sys.path.insert(0, HERE)
from hole_classifier import CAD_D, GROUP, S_PIXEL, UNKNOWN_MM
import production_plan_summary as pps

PN = json.load(open(os.path.join(HERE, 'part_numbers.json')))
BR = json.load(open(os.path.join(HERE, 'radar_bracket.json')))['classes']
OUT_MD = os.path.join(DOOR, 'DOC_dataset_spec.md'); OUT_DOCX = os.path.join(DOOR, 'DOC_dataset_spec.docx')
POS = {'LH_FRT': 'LH FRT', 'LH_RR': 'LH REAR', 'RH': 'RH'}


def q(con, sql, *a):
    return con.execute(sql, a).fetchall()


def stats(con):
    S = {}
    S['classes'] = q(con, """SELECT cl.name cls, COUNT(DISTINCT s.id) sess, SUM(i.synced_local) local_imgs,
           MIN(substr(s.session_dir,1,8)) d0, MAX(substr(s.session_dir,1,8)) d1
           FROM images i JOIN classes cl ON cl.id=i.class_id JOIN capture_sessions s ON s.id=i.session_id
           WHERE i.is_valid=1 AND cl.dataset_id=5 AND cl.name!='Unknown' GROUP BY cl.name ORDER BY cl.name""")
    S['pairs'] = {r[0]: r[1] for r in q(con, """SELECT cl.name, SUM(s.saved_pairs) FROM capture_sessions s JOIN (SELECT DISTINCT session_id, class_id FROM images WHERE is_valid=1) x ON x.session_id=s.id
           JOIN classes cl ON cl.id=x.class_id WHERE cl.name!='Unknown' GROUP BY cl.name""")}
    S['split'] = {(r[0], r[1]): r[2] for r in q(con, """SELECT cl.name, i.split, COUNT(DISTINCT i.session_id) FROM images i JOIN classes cl ON cl.id=i.class_id
           WHERE i.is_valid=1 AND i.synced_local=1 AND cl.name!='Unknown' GROUP BY 1,2""")}
    S['split_img'] = {(r[0], r[1]): r[2] for r in q(con, """SELECT cl.name, i.split, COUNT(*) FROM images i JOIN classes cl ON cl.id=i.class_id
           WHERE i.is_valid=1 AND i.synced_local=1 AND cl.name!='Unknown' GROUP BY 1,2""")}
    S['opt'] = {(r[0], r[1], r[2]): r[3] for r in q(con, """SELECT cl.name, s.option_radar, s.option_source, COUNT(DISTINCT s.id)
           FROM capture_sessions s JOIN images i ON i.session_id=s.id JOIN classes cl ON cl.id=i.class_id
           WHERE i.is_valid=1 AND cl.name!='Unknown' GROUP BY 1,2,3""")}
    S['n_sessions_all'] = q(con, "SELECT COUNT(*) FROM capture_sessions")[0][0]
    S['n_invalid'] = q(con, "SELECT COUNT(*) FROM capture_sessions s WHERE NOT EXISTS (SELECT 1 FROM images i WHERE i.session_id=s.id AND i.is_valid=1)")[0][0]
    S['n_unknown'] = q(con, "SELECT COUNT(DISTINCT s.id) FROM capture_sessions s JOIN images i ON i.session_id=s.id JOIN classes cl ON cl.id=i.class_id WHERE cl.name='Unknown'")[0][0]
    S['n_relabel'] = q(con, "SELECT COUNT(DISTINCT s.id) FROM capture_sessions s JOIN images i ON i.session_id=s.id JOIN classes cl ON cl.id=i.class_id WHERE i.is_valid=1 AND cl.name!=s.class_name AND cl.name!='Unknown'")[0][0]
    S['dates'] = q(con, "SELECT MIN(substr(session_dir,1,8)), MAX(substr(session_dir,1,8)) FROM capture_sessions s WHERE EXISTS (SELECT 1 FROM images i WHERE i.session_id=s.id AND i.is_valid=1)")[0]
    S['radar_check'] = q(con, "SELECT class_name, auto_flag, COUNT(*) FROM session_radar_check GROUP BY 1,2")
    S['radar_scores'] = q(con, "SELECT auto_flag, COUNT(*), MIN(score_med), MAX(score_med) FROM session_radar_check GROUP BY 1")
    hs = collections.defaultdict(list)
    for r in q(con, "SELECT auto_flag, params FROM session_radar_check WHERE auto_flag IS NOT NULL"):
        h = json.loads(r[1]).get('h_med') if r[1] else None
        if h is not None: hs[r[0]].append(h)
    S['radar_h'] = {k: (min(v), max(v)) for k, v in hs.items()}
    S['hole_test'] = q(con, """SELECT e.accuracy, e.correct, e.total_samples, e.evaluated_at FROM evaluation_results e WHERE e.model_id=7 AND e.per_class_results LIKE '%datasets_factory_v2/test%' ORDER BY e.evaluated_at DESC LIMIT 1""")
    S['cnn_test'] = q(con, """SELECT m.name, e.accuracy, e.correct, e.total_samples, e.evaluated_at FROM evaluation_results e JOIN models m ON m.id=e.model_id JOIN datasets d ON d.id=e.dataset_id
           WHERE e.model_id=10 AND d.name='door_factory_collect' ORDER BY e.evaluated_at DESC LIMIT 1""")
    return S


def build_md(S, plan):
    now = time.strftime('%Y-%m-%d')
    tot_sess = sum(r['sess'] for r in S['classes']); tot_local = sum(r['local_imgs'] for r in S['classes']); tot_pairs = sum(S['pairs'].values())
    L = []
    A = L.append
    A(f"# 통합모델 SIDE DOOR 현장 수집 데이터 명세서 (품번 14종)")
    A(f"")
    A(f"작성 {now} · 데이터 기준일 {S['dates'][1][:4]}-{S['dates'][1][4:6]}-{S['dates'][1][6:]} · 생성 스크립트 `partno/build_dataset_spec.py` (DB `db/door_pipeline.db` 기준, 재실행으로 갱신)")
    A("")
    A("## 1. 개요")
    A("")
    A("- **목적**: 명성공업 도어 용접 라인(경주)에서 무인 수집한 굴착기 통합모델 SIDE DOOR RGB-D 데이터를, 도어 **어셈블리 품번 14종** 기준으로 정의·집계한다. 도어 분류 AI(홀 랜드마크 판별기 + CNN 폴백)와 자세 추정의 학습·평가 데이터이며, 품번 판정 = 형상군(9) × 레이더 옵션(유/무).")
    A("- **정본 위치**: 세종 NAS `nas:Guest/weld/<날짜>/<클래스>/s_HHMMSS/` (전량). 학습 PC 로컬 미러 `datasets_factory_collect/`(세션당 20장 샘플), 학습·평가 뷰 `datasets_factory_v2/`(심볼릭 링크), 메타데이터 DB `db/door_pipeline.db`.")
    A(f"- **규모**: 유효 세션 {tot_sess} · NAS 수집 쌍 {tot_pairs:,} · 로컬 샘플 {tot_local:,}장 (기간 {S['dates'][0]}~{S['dates'][1]}, 무인 수집 정식 설치 2026-08-27). 무효 세션 {S['n_invalid']}(시험 촬영·빈 지그), 라벨 미정 {S['n_unknown']}, 라벨 정정 세션 {S['n_relabel']}.")
    A("- **주요 결정(2026-09-23)**: 품번 판정은 2단(형상군 → 레이더 옵션 → 품번 매핑), 레이더 옵션은 CAD 브래킷 투영 규칙 검사(학습 없음)로 자동 판정 후 사용자가 육안 확정, 클래스 체계(9종)와 태블릿 라벨은 유지하고 DB 에 옵션 컬럼만 추가, 공식 지표는 품번 정확도, CAD 는 필요한 것만 교체. AVM 옵션은 구분하지 않음.")
    A("")
    A("## 2. 품번 체계 (14)")
    A("")
    A("사양표·도면(2026-09-23 입수 `cad/통합모델 SIDE DOOR 도면 모음/`) 기준. 키는 리비전을 뺀 기본 품번, 리비전은 별도 필드. E23/E25 는 LH REAR·RH 를 공용하고 E30/E35/E38 은 RH 를 공용하므로 E23 고유 품번은 LH FRT 하나뿐이다. RADAR 옵션은 LH REAR·RH 에만 있다.")
    A("")
    A("| 품번 | rev | 명칭 | 적용 기종 | 위치 | 레이더 | 사양표 폭(mm) | CAD D(mm) | 파이프라인 클래스 | 도면/STEP |")
    A("|---|---|---|---|---|---|---:|---:|---|---|")
    for p in PN['parts']:
        A(f"| **{p['part_no']}** | {p['rev']} | {p['name']} | {'/'.join(p['models'])} | {POS[p['position']]} | {'O' if p['radar'] == 1 else 'X' if p['radar'] == 0 else '—'} | {p['width_mm']} | {CAD_D[p['class_name']]} | {p['class_name']} | {p['drawing'].rsplit('.', 1)[0]} |")
    A("")
    A(f"- **형상군 클래스(9)**: 좌우 세로 프레임 상단 모서리 홀(Ø6) 2개 거리 D = 폭 − 106mm 로 판정(`hole_classifier.CAD_D`, 최근접 CAD, 편차 > {UNKNOWN_MM:.0f}mm 는 unknown). D 는 힌지↔래치 픽셀 폭 × S_PIXEL/fx (S_PIXEL[54910212] = {S_PIXEL.get(54910212)} mm). 그룹: FRT {sorted(c for c in CAD_D if GROUP[c] == 'FRT')}, RR {sorted(c for c in CAD_D if GROUP[c] == 'RR')}, RH {sorted(c for c in CAD_D if GROUP[c] == 'RH')}.")
    A("- **레이더 옵션 정의**: 중앙 가로 보강재(MID) 위 브래킷 판(약 346×127mm, 6홀 평면보다 카메라 쪽 70mm)에 Ø34 홀 2개(간격 91.6mm)가 있으면 레이더 사양(1), 없으면 미장착(0). FR/RR 세로 보강재(코너 홀)·상하 보강재·힌지·볼트 사각(157×96)은 두 사양이 같은 품번이라 형상군 D 와 자세 랜드마크는 옵션과 무관. FRT 는 옵션 없음(NULL).")
    A("- **브래킷 좌표(도어 프레임, mm)**: " + "; ".join(f"{c} {v['holes']}" for c, v in sorted(BR.items())) + " (`partno/radar_bracket.json`, CAD STL 카메라 쪽 최대 z 지도에서 추출).")
    A("- 사양표 폭은 CAD D + 106 보다 일관되게 2mm 작다(코너 홀 오프셋 52/53mm 차이). 판정에는 현장 실측으로 검증된 CAD D 를 쓴다.")
    A("")
    A("## 3. 파일 구조와 형식")
    A("")
    A("```")
    A("<날짜 YYYYMMDD>/<태블릿 클래스명>/s_HHMMSS/        # 세션 = 도어 1장, 06_factory_capture.py 가 2초 간격 최대 5분 저장")
    A("  rgb_NNNN.png     1920×1200 RGB 8-bit PNG (ZED X Mini 좌안, rectified)")
    A("  depth_NNNN.png   1920×1200 16-bit PNG, 단위 mm (0 = 무효)")
    A("  meta.json        class_name, started_at/ended_at, camera, capture_interval_s, session_duration_s, intrinsics{serial,width,height,fx,fy,cx,cy,rectified}, saved_pairs, finished, stop_reason")
    A("datasets_factory_collect/  위 구조의 로컬 미러(세션당 20장 균등 샘플, db/pull_nas.py)")
    A("datasets_factory_v2/{all,train,val,test}/<DB 클래스>/rgb_<날짜>_<세션>_<idx>.png (+depth_)   # 심볼릭 링크 뷰 + manifest.json")
    A("```")
    A("")
    A("- DB(`db/schema.sql`): `capture_sessions`(세션 메타·태블릿 원본 라벨 `class_name`·이력 `notes`·**옵션 `option_radar/option_source/option_note/option_at`**), `images`(정정 라벨 `class_id`·`is_valid`·`split`·`synced_local`), `session_hole_metrics`(세션별 홀 판정 지표·다수결 `pred_major`), `session_radar_check`(레이더 자동 검사 근거), `evaluation_results`.")
    A("- 카메라: ZED X Mini(SN 54910212) 1920×1200, fx=fy≈1269.7, 고정 지그 위 약 1.45m, 도어 안쪽 면(보강재 면)을 본다. 실측 intrinsics 는 meta.json 에 기록되며 판정은 픽셀 폭 기준이라 depth 스케일 변동(세션 ±1.4%)에 무관.")
    A("")
    A("## 4. 라벨 정의")
    A("")
    A("- **형상군 라벨**: 태블릿(06) 입력 8종(E23 없음)이 원본(`capture_sessions.class_name`, 불변). DB 정정 라벨(`images.class_id`)이 정본이며 정정 이력은 `notes` 에 남는다. 정정 규칙: 홀 판별기 세션 다수결이 라벨과 다르고 두 클래스 CAD D 차 ≥100mm·판정 ≥10장·일치 ≥90% 이면 자동 재배정(`db/build_dataset.py auto-relabel`, E23↔E25 FRT 등), 인접 쌍(E30↔E38 FRT 47mm)은 육안 확인.")
    A("- **레이더 옵션 라벨(세션 단위)**: `partno/radar_check.py` 가 프레임별로 브래킷 홀 2개 투영 위치의 어두운 홀 유무를 점수화(≥0.30 있음 / ≤0.12 없음)하고 세션 표(판정 ≥5장·다수 ≥80%)로 자동 플래그(`option_source='auto'`). 사용자가 웹 도구 `/options` 에서 크롭 몽타주를 보고 확정하면 `option_source='user'`(자동 재검사가 덮어쓰지 않음). 품번 정확도의 GT 는 사용자 확정값만 쓴다.")
    A("- **무효화**: 촬영 조건(설치 전 시험 촬영, 빈 지그, 작업자 유입)이 이유일 때만 `invalidate`. 판정이 틀렸다는 이유로는 무효화하지 않는다.")
    A("- **홀 랜드마크 라벨**(검출기 학습용, 별도): `labels/holes/*.json` 134장(볼트 4·코너 2 점 좌표, `tools/label_holes.py`).")
    A("- **분할**: 세션 단위만(프레임 단위 금지). 클래스별 시간순 test=max(1,15%) → val → train, 가장 이른 세션이 test 로 고정(현장 벤치마크), 세션 <3 클래스는 전부 train(평가 제외).")
    A("")
    A("## 5. 수량")
    A("")
    A("### 5.1 형상군 클래스별")
    A("")
    A("| 클래스 | 유효 세션 | NAS 쌍 | 로컬 샘플 | 기간 | test/val/train 세션 | test 장수 |")
    A("|---|---:|---:|---:|---|---|---:|")
    for r in S['classes']:
        sp = lambda k: S['split'].get((r['cls'], k), 0)
        A(f"| {r['cls']} | {r['sess']} | {S['pairs'].get(r['cls'], 0):,} | {r['local_imgs']} | {r['d0']}~{r['d1']} | {sp('test')}/{sp('val')}/{sp('train')}{(' (+미지정 ' + str(sp(None)) + ')') if sp(None) else ''} | {S['split_img'].get((r['cls'], 'test'), 0)} |")
    A(f"| **계** | **{tot_sess}** | **{tot_pairs:,}** | **{tot_local:,}** | {S['dates'][0]}~{S['dates'][1]} | | |")
    A("")
    A("### 5.2 품번(옵션)별 세션 — 레이더 확정값(user)과 자동 판정(auto) 구분")
    A("")
    A("| 클래스 | 품번(레이더 X) | X 확정 | X 자동 | 품번(레이더 O) | O 확정 | O 자동 | 미정·미검사 |")
    A("|---|---|---:|---:|---|---:|---:|---:|")
    pn_by = {(p['class_name'], p['radar']): p['part_no'] for p in PN['parts']}
    for r in S['classes']:
        c = r['cls']; o = lambda v, src: S['opt'].get((c, v, src), 0)
        if (c, None) in pn_by:
            A(f"| {c} | {pn_by[(c, None)]} (옵션 없음) | | | | | | {r['sess']} |")
        else:
            A(f"| {c} | {pn_by[(c, 0)]} | {o(0, 'user')} | {o(0, 'auto')} | {pn_by[(c, 1)]} | {o(1, 'user')} | {o(1, 'auto')} | {o(None, None) + o(None, 'auto')} |")
    sc = {r[0]: r for r in S['radar_scores']}; hr = S['radar_h']
    A("")
    A(f"자동 검사 두 단서의 세션 중앙값 분포 — 레이더 O {sc.get(1, [0, 0, 0, 0])[1]}세션: 밝기 {sc.get(1, [0,0,0,0])[2] or 0:.2f}~{sc.get(1, [0,0,0,0])[3] or 0:.2f}, 판 높이 {hr.get(1, (0, 0))[0]:.0f}~{hr.get(1, (0, 0))[1]:.0f}mm · "
      f"X {sc.get(0, [0,0,0,0])[1]}세션: 밝기 {sc.get(0, [0,0,0,0])[2] or 0:.2f}~{sc.get(0, [0,0,0,0])[3] or 0:.2f}, 판 높이 {hr.get(0, (0, 0))[0]:.0f}~{hr.get(0, (0, 0))[1]:.0f}mm · "
      f"미정 {sc.get(None, [0,0,0,0])[1]}세션(프레임 1~3장·빈 지그). 판 높이 단서에서 두 사양이 겹치지 않는다(브래킷 판 CAD 70mm vs 보강재 29mm, depth 는 가장자리를 뭉개 낮게 읽힘).")
    A("")
    A("### 5.3 생산계획 대조 (월 단위 비율, 1:1 매칭 아님)")
    A("")
    A("생산계획(`cad/9-10월_통합모델_일별생산계획.xlsx`, 고객사 L/ON 일자 기준 210대)의 레이더 표기('레이더 X'=미장착, 무표기=장착, '확인하기'=미정)로 집계한 기종별 장착 비율. 도어 제작은 L/ON 보다 1주 이상 앞서고 카메라는 도어의 일부(주별 40~60%)만 잡으므로 세션↔개체 매칭은 하지 않는다.")
    A("")
    A("| 월 | 그룹 | 개체 | 레이더 장착 | 미장착 | 미정 | 장착 비율 |")
    A("|---|---|---:|---:|---:|---:|---:|")
    for (mon, g), d in sorted(plan.items()):
        r1, r0, rn = d['radar'][1], d['radar'][0], d['radar'][None]
        A(f"| {mon} | {g} | {d['units']} | {r1} | {r0} | {rn} | {100 * r1 / max(1, r1 + r0):.0f}% |")
    A("")
    A("현장 자동 판정의 레이더 비율(RR/RH 세션)은 계획 비율보다 높다(예: 9월 E23/E25 RR·RH 40%대 vs 계획 15%). 카메라가 잡은 도어의 편중 또는 표기 해석 차이일 수 있어 공장 확인 항목으로 둔다.")
    A("")
    A("## 6. 평가 정의와 현재 수치")
    A("")
    A("- **판정률** = 품번이 나온 세션(프레임) 비율(형상군 판정 + RR/RH 면 레이더 판정 모두 있음). **판정 정확도** = 판정된 것 중 정답 비율. 공식 수치는 고정 test 세션(형상군 벤치마크와 동일 세션), 전량은 참고(`partno/evaluate_partno.py`).")
    A("- 품번 GT = DB 정정 클래스 × 사용자 확정 레이더. 확정 전 세션은 분모에서 제외하고 수를 적는다. 2026-09-23 확정은 자동 판정 145세션의 브래킷 크롭을 사용자가 세션별로 육안 검토해 일괄 승인한 것이므로, 레이더 항목의 정확도는 '규칙 판정과 사람 검토의 일치율'이며 블라인드 시험이 아니다. 독립 검증은 이후 유입 세션을 확정 전에 판정해 비교하는 방식으로 한다.")
    ht = S['hole_test'][0] if S['hole_test'] else None; ct = S['cnn_test'][0] if S['cnn_test'] else None
    if ht: A(f"- 형상군(홀 판별기) 고정 test: {ht[1]}/{ht[2]} = {ht[0]:.1f}% ({ht[3][:10]}).")
    if ct: A(f"- CNN 폴백({ct[0]}) 고정 test: {ct[2]}/{ct[3]} = {ct[1]:.1f}% ({ct[4][:10]}).")
    ev = os.path.join(HERE, 'eval_partno.json')
    if os.path.exists(ev):
        e = json.load(open(ev))
        for k, lab in (('session_test', '품번 세션 단위 · 고정 test'), ('frame_test', '품번 프레임 단위 · 고정 test'), ('session_all', '품번 세션 단위 · 전량')):
            r = e.get(k)
            if r: A(f"- {lab}: 대상 {r['n']}(GT 확정 {r['n_gt']}, 미확정 {r['n_gt_missing']}) · 판정률 {r['judged_rate']:.1f}% · 판정 정확도 {r['acc_judged']:.1f}% ({r['correct']}/{r['judged']}).")
    else:
        A("- 품번 정확도: 사용자 레이더 확정 후 `partno/evaluate_partno.py` 실행 시 기록.")
    rp = os.path.normpath(os.path.join(DOOR, '..', 'door_partno', 'report'))
    bl = os.path.join(rp, 'blind_unsplit_detblind.json'); cvj = os.path.join(rp, 'cv_cvA_multi_448.json')
    if os.path.exists(bl):
        b = json.load(open(bl)); R = b['results']
        A("")
        A("### 6.1 학습형 판정기 비교 (door_partno, 2026-09-23)")
        A("")
        A(f"확정 라벨로 4채널 홀 랜드마크 검출기(브래킷 채널 추가)와 CNN(9클래스+레이더 헤드 / 14클래스 / 640 변형)을 학습해, 규칙 검사와 같은 프레임에서 비교했다. 블라인드 집합 = 어느 모델도 학습에 쓰지 않은 {len(b['sessions'])}세션(9/16 오후~9/23 유입).")
        A("")
        A("| 판정기 | 블라인드 판정률 | 블라인드 품번 정확도 | 세션 |")
        A("|---|---:|---:|---:|")
        names = dict(rule='규칙 검사(운영)', det='4채널 검출기(블라인드 세션 제외 학습)', partno_multi_448_seed42='CNN 9클래스+레이더 헤드 448', partno_part14_448_seed42='CNN 14클래스 448', partno_multi_640_seed42='CNN 9클래스+레이더 헤드 640')
        for k, v in R.items():
            A(f"| {names.get(k, k)} | {v['judged_rate']:.1f}% | {v['acc']:.1f}% ({v['correct']}/{v['judged']}) | {v['s_correct']}/{v['sessions']} |")
        cvtxt = ''
        if os.path.exists(cvj):
            c = json.load(open(cvj)); tot = sum(x['part_n'] for x in c.values()); ok = sum(x['part_ok'] for x in c.values())
            cvtxt = f" CNN 9클래스+레이더 헤드 448 의 세션 5-fold 교차검증은 {ok}/{tot} = {100 * ok / max(1, tot):.1f}%."
        A("")
        A(f"- CNN 오류는 전부 형상군 혼동(E30↔E38 FRT 47mm 인접 쌍 등)이고 레이더 오류는 0 이었다.{cvtxt} 결론: 형상군은 홀 거리 규칙, 레이더는 검출기 브래킷 채널이 새 세션에서도 100% 를 유지 → 4채널 검출기를 배포 정본으로 채택하고 실시간 서버에는 브래킷 피크 → 규칙 하이브리드로 통합했다(2026-09-23).")
    A("")
    A("## 7. 도구")
    A("")
    A("`partno/README.md` 의 루틴: ingest/pull → `20_session_drift.py`(다수결) → `auto-relabel` → `partno/radar_check.py` → 웹 `/options` 확정 → auto-split/build → `17_evaluate_hole_classifier.py` → `partno/evaluate_partno.py` → `partno/production_plan_summary.py`. 임의 폴더 추론은 `partno/infer_partno.py`.")
    A("")
    A("## 8. 한계·미확인")
    A("")
    A("- AVM 옵션은 구분하지 않는다(사용자 결정 2026-09-23). 생산계획의 '확인하기' 표기 개체는 미정. 레이더 표기 해석(X=미장착)은 공장 확인 전.")
    A("- 레이더 검사는 브래킷이 클램프·손에 가리면 미정. 카메라 관측면·거치 거리가 바뀌면 시차 보정값(z_plate 70mm)·탐색 폭 재확인 필요. 새 리비전에서 브래킷 위치가 바뀌면 `extract_radar_bracket.py` 재실행.")
    A("- E23 자세 CAD 는 어셈블리 STEP(110982-02444B) 메시로 2026-09-23 정식 등록(코너 홀 거리 457.6mm → CAD_D 458, 현장 정합 3.1mm, 잠정 등록 해제). E30_door_LH_RR 는 종전 메시(구 리비전, RADAR 사양) 유지 — 신규 02360K 는 홀 추출이 되는 세밀 메시(약 420MB) 재변환이 필요해 보류(`cad/door_stl/README_mesh_20260923.md`).")
    A("- 카메라는 라인 도어의 일부만 촬영하므로 수량 비율은 생산 비율과 다를 수 있다. 명성 자체 작업 실적(품번·시각)을 받으면 세션 1:1 GT 로 확장 가능.")
    return "\n".join(L) + "\n"


def md_to_docx(md, path):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.oxml.ns import qn
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']; st.font.name = 'Malgun Gothic'; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')

    def add_runs(par, text, mono=False):
        for tok in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
            if not tok: continue
            if tok.startswith('**'): r = par.add_run(tok[2:-2]); r.bold = True
            elif tok.startswith('`'): r = par.add_run(tok[1:-1]); r.font.name = 'Consolas'
            else: r = par.add_run(tok)
            if mono: r.font.name = 'Consolas'; r.font.size = Pt(8.5)
    lines = md.splitlines(); i = 0; table = []
    def flush_table():
        nonlocal table
        rows = [[c.strip() for c in ln.strip().strip('|').split('|')] for ln in table if not re.match(r'^\|\s*-', ln)]
        if rows:
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row[:len(rows[0])]):
                    p = t.cell(ri, ci).paragraphs[0]; add_runs(p, cell)
                    for r in p.runs: r.font.size = Pt(8); r.bold = r.bold or ri == 0
        table = []
    while i < len(lines):
        ln = lines[i]
        if ln.startswith('|'):
            table.append(ln); i += 1; continue
        if table: flush_table()
        if ln.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                add_runs(doc.add_paragraph(), lines[i], mono=True); i += 1
            i += 1; continue
        m = re.match(r'^(#+)\s+(.*)', ln)
        if m: doc.add_heading(m.group(2), level=min(len(m.group(1)), 3)); i += 1; continue
        if ln.startswith('- '): add_runs(doc.add_paragraph(style='List Bullet'), ln[2:]); i += 1; continue
        if ln.strip(): add_runs(doc.add_paragraph(), ln)
        i += 1
    if table: flush_table()
    doc.save(path)


if __name__ == '__main__':
    ap = argparse.ArgumentParser(); ap.add_argument('--db', default=os.path.join(DOOR, 'db', 'door_pipeline.db')); ap.add_argument('--no-docx', action='store_true')
    a = ap.parse_args()
    con = sqlite3.connect(a.db); con.row_factory = sqlite3.Row
    S = stats(con)
    _, units = pps.read_units(pps.DEFAULT_XLSX); plan, _ = pps.summarize(units)
    md = build_md(S, plan)
    open(OUT_MD, 'w', encoding='utf8').write(md); print(f"→ {OUT_MD} ({len(md.splitlines())}줄)")
    if not a.no_docx:
        md_to_docx(md, OUT_DOCX); print(f"→ {OUT_DOCX}")
